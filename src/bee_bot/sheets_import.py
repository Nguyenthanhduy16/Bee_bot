"""Read an XLSX export into a *local*, provisional schedule snapshot.

This does not synchronize Sheets, establish user identity, prove programme-wide
coverage, or create production assignment IDs. Source exports are never modified.
"""
from collections import Counter, defaultdict
from datetime import date, datetime, time
from hashlib import sha256
import json
import os
from pathlib import Path
import re
import tempfile
import unicodedata
from zoneinfo import ZoneInfo

import openpyxl

from .errors import BeeBotError

ZONE = ZoneInfo('Asia/Ho_Chi_Minh')
MERGED_COLUMNS = {2, 3, 5, 6, 7, 8}
DATE_PATTERN = re.compile(r'(?<!\d)(\d{1,2})/(\d{1,2})/(\d{4})(?!\d)')


def clean(value):
    return ' '.join(unicodedata.normalize('NFC', str(value or '')).split())


def normalized(value):
    return clean(value).casefold()


def digest(prefix, *parts):
    serialized = json.dumps(parts, ensure_ascii=False, separators=(',', ':'))
    return prefix + sha256(serialized.encode()).hexdigest()[:20]


def _date_header(value):
    """Reject workbook title/ranges, even when they contain a complete date."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean(value)
    matches = list(DATE_PATTERN.finditer(text))
    if len(matches) != 1:
        return None
    match = matches[0]
    prefix, suffix = text[:match.start()].strip(' -–—|:'), text[match.end():].strip()
    if suffix or (prefix and not re.fullmatch(r'(?:thứ\s+(?:\d+|hai|ba|tư|năm|sáu|bảy)|chủ\s+nhật|ngày)', prefix, re.I)):
        return None
    try:
        return date(int(match[3]), int(match[2]), int(match[1]))
    except ValueError:
        return None


def _numeric_row(value):
    return (isinstance(value, (int, float)) and not isinstance(value, bool)
            and value > 0 and float(value).is_integer())


def import_xlsx(path):
    source = Path(path)
    try:
        workbook = openpyxl.load_workbook(source, data_only=True, read_only=False)
    except (OSError, ValueError, KeyError) as exc:
        raise BeeBotError('IMPORT_FAILED', 'Không đọc được bản xuất XLSX.') from exc
    people, assignments, warnings = {}, [], []
    stats = Counter(numeric_rows=0, imported_positions=0, assigned_positions=0,
                    empty_positions=0, skipped_numeric_rows=0, repeated_region_mismatches=0)
    days_with_rows, observed_days = Counter(), set()
    occurrences = Counter()
    name_variants = defaultdict(set)
    sheet_titles = []
    try:
        for sheet in workbook:
            sheet_titles.append(sheet.title)
            anchors = {}
            for merged in sheet.merged_cells.ranges:
                for column in MERGED_COLUMNS:
                    if merged.min_col <= column <= merged.max_col:
                        for row in range(merged.min_row, merged.max_row + 1):
                            anchors[(row, column)] = (merged.min_row, merged.min_col)
            current_day, current_day_row = None, None
            for row in range(1, sheet.max_row + 1):
                a = sheet.cell(row, 1).value
                header = _date_header(a)
                if header:
                    current_day, current_day_row = header, row
                    observed_days.add((sheet.title, header.isoformat()))
                    continue
                if isinstance(a, str) and re.match(r'^(?:thứ\b|chủ\s+nhật\b|ngày\b)', clean(a), re.I) and DATE_PATTERN.search(a):
                    current_day, current_day_row = None, None
                    warnings.append({'code': 'INVALID_DATE_HEADER', 'sheet': sheet.title, 'row': row})
                if not _numeric_row(a):
                    continue
                stats['numeric_rows'] += 1

                def value(column):
                    anchor = anchors.get((row, column))
                    # A malformed merge across day blocks must not carry old values.
                    if anchor:
                        if current_day_row and anchor[0] <= current_day_row:
                            return None
                        return sheet.cell(*anchor).value
                    return sheet.cell(row, column).value

                track, session = clean(value(2)), clean(value(3))
                room, daypart = clean(value(6)), normalized(value(7))
                times = {'sáng': (9, 13), 'chiều': (14, 18)}.get(daypart)
                missing = [name for name, exists in (
                    ('date', current_day is not None), ('track', bool(track)),
                    ('session', bool(session)), ('room', bool(room)), ('daypart', times is not None)) if not exists]
                if missing:
                    stats['skipped_numeric_rows'] += 1
                    warnings.append({'code': 'INVALID_CANONICAL_ROW', 'sheet': sheet.title,
                                     'row': row, 'missing_fields': missing})
                    continue
                days_with_rows[(sheet.title, current_day.isoformat())] += 1
                name = clean(sheet.cell(row, 4).value)  # Never carry D or infer from J.
                coach_id = None
                if name:
                    name_key = normalized(name)
                    name_variants[name_key].add(name)
                    coach_id = digest('source_person_', name_key)
                    people.setdefault(coach_id, {'person_id': coach_id, 'display_name': name,
                        'role': 'labcoach', 'active': False, 'identity_verified': False,
                        'identity_source': 'provisional_name_only'})
                    stats['assigned_positions'] += 1
                else:
                    stats['empty_positions'] += 1
                class_key = (sheet.title, current_day.isoformat(), normalized(track),
                             normalized(session), normalized(room), daypart)
                shift_id = digest('source_shift_', *class_key)
                occurrences[class_key] += 1
                assignment_id = digest('source_assignment_', *class_key, occurrences[class_key])
                key_coach = clean(value(5))
                repeated = {11: current_day.isoformat(), 12: track, 13: session,
                            14: room, 15: daypart, 16: key_coach}
                differing = []
                for column, canonical in repeated.items():
                    other = sheet.cell(row, column).value
                    if column == 11:
                        other_day = _date_header(other)
                        equal = other_day is not None and other_day.isoformat() == canonical
                    else:
                        equal = normalized(other) == normalized(canonical)
                    # Blank repeated cells are diagnostics only; never fill canonical data.
                    if clean(other) and not equal:
                        differing.append(openpyxl.utils.get_column_letter(column))
                if differing:
                    stats['repeated_region_mismatches'] += 1
                    warnings.append({'code': 'REPEATED_REGION_MISMATCH', 'sheet': sheet.title,
                                     'row': row, 'columns': differing})
                assignment = {'assignment_id': assignment_id, 'shift_id': shift_id,
                    'coach_id': coach_id, 'starts_at': datetime.combine(current_day, time(times[0]), ZONE).isoformat(),
                    'ends_at': datetime.combine(current_day, time(times[1]), ZONE).isoformat(),
                    'room': room, 'site': 'unverified', 'class_name': f'{track} / {session}',
                    'track': track, 'session': session, 'daypart': clean(value(7)),
                    'key_coach': key_coach, 'room_capacity': value(8),
                    'j': sheet.cell(row, 10).value, 'revision': 0,
                    'source_locator': {'sheet': sheet.title, 'row': row, 'assignment_cell': f'D{row}',
                                       'availability_cell': f'J{row}'},
                    'provenance': {'format': 'xlsx_export', 'canonical_region': 'B:J',
                                   'identity_verified': False, 'id_scope': 'snapshot_only'},
                    'slot_occurrence': occurrences[class_key]}
                assignments.append(assignment)
                stats['imported_positions'] += 1
        for sheet, day in sorted(observed_days):
            if not days_with_rows[(sheet, day)]:
                warnings.append({'code': 'DATE_WITHOUT_NUMERIC_ROWS', 'sheet': sheet, 'date': day})
        warnings.append({'code': 'UNVERIFIED_NAME_IDENTITIES', 'count': len(people),
            'detail': 'Name-only groups cannot distinguish homonymous people or verify accounts.'})
        variant_count = sum(len(variants) > 1 for variants in name_variants.values())
        if variant_count:
            warnings.append({'code': 'NAME_VARIANTS_REQUIRE_REVIEW', 'count': variant_count})
        people['source_ops'] = {'person_id': 'source_ops', 'display_name': 'Vận hành (chưa xác thực)',
                               'role': 'ops', 'active': False, 'identity_verified': False,
                               'identity_source': 'synthetic_placeholder'}
        stats['provisional_labcoach_groups'] = len(people) - 1
        stats['observed_date_headers'] = len(observed_days)
        metadata = {'source_kind': 'read_only_xlsx_export', 'source_filename': source.name,
            'source_sha256': sha256(source.read_bytes()).hexdigest(), 'sheet_titles': sheet_titles,
            'imported_at': datetime.now(ZONE).isoformat(), 'timezone': str(ZONE),
            'complete_coverage': False, 'requires_verified_identity': True,
            'assignment_id_scope': 'snapshot_only', 'source_write_enabled': False,
            'stats': dict(stats), 'warnings': warnings}
        return {'people': list(people.values()), 'assignments': assignments, 'metadata': metadata}
    finally:
        workbook.close()


def write_snapshot(input_path, output_path):
    input_file, output_file = Path(input_path), Path(output_path)
    if (input_file.resolve() == output_file.resolve()
            or (output_file.exists() and input_file.samefile(output_file))):
        raise BeeBotError('SOURCE_OVERWRITE_FORBIDDEN', 'Không được ghi đè bản xuất nguồn.')
    snapshot = import_xlsx(input_file)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile('w', dir=output_file.parent, encoding='utf-8', delete=False) as handle:
            temporary = handle.name
            json.dump(snapshot, handle, ensure_ascii=False, indent=2)
            handle.write('\n')
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, output_file)
    finally:
        if temporary and os.path.exists(temporary):
            os.unlink(temporary)
    metadata = snapshot['metadata']
    return {**metadata['stats'], 'output_path': str(output_file),
            'complete_coverage': False, 'requires_verified_identity': True,
            'warning_counts': dict(Counter(w['code'] for w in metadata['warnings']))}
