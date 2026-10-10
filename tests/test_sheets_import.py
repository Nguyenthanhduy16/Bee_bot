"""Synthetic workbooks only; no source names or private XLSX fixtures committed."""
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from bee_bot.errors import BeeBotError
from bee_bot.sheets_import import import_xlsx, write_snapshot


def workbook(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Lịch thử'
    sheet['A1'] = 'LỊCH LAB | 02/10–16/10/2026'
    sheet['A2'] = 1  # Must not inherit the title's date.
    sheet['B2'], sheet['C2'], sheet['F2'], sheet['G2'] = 'T0', 'CS0', 'H000', 'Sáng'
    sheet['A3'] = 'Thứ 2 - 12/10/2026'
    for row in (4, 5):
        sheet.cell(row, 1, row - 3)
        sheet.cell(row, 11, '12/10/2026')
        for column, value in ((12, 'Track 1'), (13, 'CS1'), (14, 'H210'), (15, 'Sáng'), (16, 'Key')):
            sheet.cell(row, column, value)
    for column, value in ((2, 'Track 1'), (3, 'CS1'), (5, 'Key'), (6, 'H210'), (7, 'Sáng'), (8, 30)):
        sheet.cell(4, column, value)
        sheet.merge_cells(start_row=4, start_column=column, end_row=5, end_column=column)
    sheet['D4'], sheet['J4'] = 'Người thử', 'Bận'
    sheet['J5'] = 'Rảnh'
    sheet['N5'] = 'H205'  # F remains canonical.
    sheet['A6'] = 'Thứ 3 - 13/10/2026'
    sheet['A7'], sheet['D7'], sheet['J7'] = 1, 'Người thử', 'Rảnh'
    sheet['K7'], sheet['L7'], sheet['M7'], sheet['N7'], sheet['O7'] = '13/10/2026', 'Track 2', 'CS2', 'H205', 'Chiều'
    sheet['A8'] = 'Thứ 4 - 14/10/2026'
    path = tmp_path / 'source.xlsx'
    book.save(path)
    book.close()
    return path


def test_canonical_merged_scope_empty_d_and_title_date(tmp_path):
    path = workbook(tmp_path)
    snapshot = import_xlsx(path)
    rows = snapshot['assignments']
    assert len(rows) == 2
    assert rows[0]['coach_id']
    assert rows[1]['coach_id'] is None
    assert rows[0]['j'] == 'Bận' and rows[1]['j'] == 'Rảnh'
    assert all(row['room'] == 'H210' for row in rows)
    assert all(row['starts_at'] == '2026-10-12T09:00:00+07:00' for row in rows)
    assert all(row['ends_at'] == '2026-10-12T13:00:00+07:00' for row in rows)
    assert rows[0]['shift_id'] == rows[1]['shift_id']
    assert rows[0]['assignment_id'] != rows[1]['assignment_id']
    warnings = snapshot['metadata']['warnings']
    assert any(w['code'] == 'REPEATED_REGION_MISMATCH' and w['row'] == 5 for w in warnings)
    assert any(w['code'] == 'INVALID_CANONICAL_ROW' and w['row'] == 7 for w in warnings)
    assert snapshot['metadata']['stats']['skipped_numeric_rows'] == 2


def test_import_flags_do_not_verify_names_coverage_or_production_ids(tmp_path):
    snapshot = import_xlsx(workbook(tmp_path))
    assert snapshot['metadata']['complete_coverage'] is False
    assert snapshot['metadata']['requires_verified_identity'] is True
    assert snapshot['metadata']['source_write_enabled'] is False
    assert snapshot['metadata']['assignment_id_scope'] == 'snapshot_only'
    assert all(person['active'] is False and person['identity_verified'] is False for person in snapshot['people'])
    assert any(w['code'] == 'UNVERIFIED_NAME_IDENTITIES' for w in snapshot['metadata']['warnings'])
    assert any(w['code'] == 'DATE_WITHOUT_NUMERIC_ROWS' and w['date'] == '2026-10-14'
               for w in snapshot['metadata']['warnings'])


def test_snapshot_identifiers_are_not_row_coordinates(tmp_path):
    path = workbook(tmp_path)
    before = import_xlsx(path)
    book = openpyxl.load_workbook(path)
    sheet = book.active
    # Insert a blank row before the data using reconstruction to keep merges coherent.
    for merged in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(merged))
    for row in (4, 5):
        for column in (2, 3, 5, 6, 7, 8):
            sheet.cell(row, column, sheet.cell(4, column).value)
    sheet.insert_rows(3)
    book.save(path)
    book.close()
    after = import_xlsx(path)
    assert [a['assignment_id'] for a in before['assignments']] == [a['assignment_id'] for a in after['assignments']]
    assert before['assignments'][0]['source_locator']['row'] + 1 == after['assignments'][0]['source_locator']['row']


def test_write_snapshot_preserves_source_and_summary_contains_no_names(tmp_path):
    path = workbook(tmp_path)
    original = path.read_bytes()
    summary = write_snapshot(path, tmp_path / 'snapshot.json')
    assert path.read_bytes() == original
    assert summary['imported_positions'] == 2
    assert summary['assigned_positions'] == 1
    assert summary['empty_positions'] == 1
    assert 'Người thử' not in str(summary)
    with pytest.raises(BeeBotError):
        write_snapshot(path, path)
    assert path.read_bytes() == original


def test_datetime_header_and_afternoon_times(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append([datetime(2026, 10, 12)])
    sheet.append([1, 'Track 1', 'CS1', 'Người thử', 'Key', 'H210', 'Chiều', 30, None, 'Chưa cập nhật'])
    path = tmp_path / 'afternoon.xlsx'
    book.save(path)
    result = import_xlsx(path)
    assert result['assignments'][0]['starts_at'] == '2026-10-12T14:00:00+07:00'
    assert result['assignments'][0]['ends_at'] == '2026-10-12T18:00:00+07:00'
